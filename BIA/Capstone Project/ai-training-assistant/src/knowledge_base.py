"""Knowledge base: loads documents, chunks, embeds, and stores as vectors.

Uses ONNX Runtime for embeddings instead of PyTorch/sentence-transformers,
keeping the install lightweight (~200MB vs ~2.5GB) and avoiding Windows
path length issues.

Two interchangeable vector-store backends are supported behind a common
interface (``add`` / ``query`` / ``count`` / ``save`` / ``load``):

- ``chroma`` (default): a real, persistent local vector database (ChromaDB).
- ``numpy``: the legacy in-memory NumPy + cosine-similarity fallback.

The backend is selected via the ``VECTOR_BACKEND`` environment variable so the
rest of the app (assistant, evaluation, UI) is unchanged.
"""

import json
import os
from pathlib import Path

import numpy as np
import onnxruntime as ort
from huggingface_hub import hf_hub_download
from tokenizers import Tokenizer


DATA_DIR = Path(__file__).resolve().parents[1] / "data" / "documents"
STORE_DIR = Path(__file__).resolve().parents[1] / "data" / "vector_store"
CHROMA_DIR = Path(__file__).resolve().parents[1] / "data" / "chroma_store"
MODEL_DIR = Path(__file__).resolve().parents[1] / "data" / "model"

# ChromaDB collection name for the training-assistant knowledge base.
CHROMA_COLLECTION = "training_assistant"

# Maps a top-level corpus folder to the router's category. Chunks are tagged
# with these so retrieval can be scoped by routing category (metadata filter).
# The router classifies queries into: company / policy / onboarding / general.
FOLDER_CATEGORY_MAP = {
    "company": "company",
    "roles": "company",
    "policies": "policy",
    "admin": "onboarding",
    "faq": "onboarding",
}


def category_for_source(source: str) -> str:
    """Map a document source path to a routing category.

    Uses the first path segment of the source (e.g. ``corpus/policies/...``)
    to derive the category. Sources that do not match a known folder default
    to ``"general"``.

    Args:
        source: The relative source path, e.g. ``corpus/policies/leave.md``.

    Returns:
        One of ``"company"``, ``"policy"``, ``"onboarding"``, or ``"general"``.
    """
    parts = source.replace("\\", "/").split("/")
    # Skip a leading "corpus" segment if present, then read the group folder.
    segments = [p for p in parts if p and p != "corpus"]
    group = segments[0].lower() if len(segments) > 1 else ""
    return FOLDER_CATEGORY_MAP.get(group, "general")


def get_backend() -> str:
    """Return the configured vector-store backend name.

    Reads the ``VECTOR_BACKEND`` environment variable and normalizes it.
    Unknown values fall back to ``"chroma"``.

    Returns:
        Either ``"chroma"`` or ``"numpy"``.
    """
    backend = os.getenv("VECTOR_BACKEND", "chroma").strip().lower()
    return backend if backend in {"chroma", "numpy"} else "chroma"

# all-MiniLM-L6-v2: small, fast, 384-dim embeddings
MODEL_REPO = "sentence-transformers/all-MiniLM-L6-v2"
ONNX_FILE = "onnx/model.onnx"
TOKENIZER_FILE = "tokenizer.json"


class TextSplitter:
    """Simple recursive text splitter."""

    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 100):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.separators = ["\n\n", "\n", ". ", " "]

    def split_text(self, text: str) -> list[str]:
        """Split text into overlapping chunks."""
        return self._split_recursive(text, self.separators)

    def _split_recursive(self, text: str, separators: list[str]) -> list[str]:
        if len(text) <= self.chunk_size:
            return [text.strip()] if text.strip() else []

        separator = separators[0] if separators else ""
        for sep in separators:
            if sep in text:
                separator = sep
                break

        parts = text.split(separator) if separator else [text]
        chunks: list[str] = []
        current_chunk = ""

        for part in parts:
            candidate = (current_chunk + separator + part) if current_chunk else part
            if len(candidate) <= self.chunk_size:
                current_chunk = candidate
            else:
                if current_chunk.strip():
                    chunks.append(current_chunk.strip())
                if len(part) > self.chunk_size and len(separators) > 1:
                    sub_chunks = self._split_recursive(part, separators[1:])
                    chunks.extend(sub_chunks)
                    current_chunk = ""
                else:
                    current_chunk = part

        if current_chunk.strip():
            chunks.append(current_chunk.strip())

        if self.chunk_overlap > 0 and len(chunks) > 1:
            overlapped: list[str] = [chunks[0]]
            for i in range(1, len(chunks)):
                prev_tail = chunks[i - 1][-self.chunk_overlap:]
                overlapped.append(prev_tail + " " + chunks[i])
            return overlapped

        return chunks


class EmbeddingModel:
    """Lightweight embedding model using ONNX Runtime (no PyTorch needed).

    Downloads and caches the all-MiniLM-L6-v2 ONNX model from HuggingFace.
    """

    def __init__(self):
        MODEL_DIR.mkdir(parents=True, exist_ok=True)

        # Download model files if not cached
        onnx_path = MODEL_DIR / "model.onnx"
        tokenizer_path = MODEL_DIR / "tokenizer.json"

        if not onnx_path.exists():
            print("📥 Downloading embedding model (first run only)...")
            downloaded = hf_hub_download(
                repo_id=MODEL_REPO,
                filename=ONNX_FILE,
                local_dir=MODEL_DIR,
                local_dir_use_symlinks=False,
            )
            # Move to expected location if nested
            src = Path(downloaded)
            if src != onnx_path:
                onnx_path.write_bytes(src.read_bytes())

        if not tokenizer_path.exists():
            downloaded = hf_hub_download(
                repo_id=MODEL_REPO,
                filename=TOKENIZER_FILE,
                local_dir=MODEL_DIR,
                local_dir_use_symlinks=False,
            )
            src = Path(downloaded)
            if src != tokenizer_path:
                tokenizer_path.write_bytes(src.read_bytes())

        # Load tokenizer and ONNX model
        self.tokenizer = Tokenizer.from_file(str(tokenizer_path))
        self.tokenizer.enable_truncation(max_length=256)
        self.tokenizer.enable_padding(length=256)

        self.session = ort.InferenceSession(
            str(onnx_path),
            providers=["CPUExecutionProvider"],
        )

    def _mean_pooling(self, token_embeddings: np.ndarray, attention_mask: np.ndarray) -> np.ndarray:
        """Apply mean pooling to token embeddings."""
        mask_expanded = np.expand_dims(attention_mask, axis=-1)
        sum_embeddings = np.sum(token_embeddings * mask_expanded, axis=1)
        sum_mask = np.sum(mask_expanded, axis=1)
        return sum_embeddings / np.maximum(sum_mask, 1e-9)

    def _normalize(self, embeddings: np.ndarray) -> np.ndarray:
        """L2 normalize embeddings."""
        norms = np.linalg.norm(embeddings, axis=1, keepdims=True)
        return embeddings / np.maximum(norms, 1e-9)

    def encode(self, texts: list[str]) -> np.ndarray:
        """Encode texts into normalized embeddings."""
        encoded = self.tokenizer.encode_batch(texts)

        input_ids = np.array([e.ids for e in encoded], dtype=np.int64)
        attention_mask = np.array([e.attention_mask for e in encoded], dtype=np.int64)
        token_type_ids = np.zeros_like(input_ids, dtype=np.int64)

        outputs = self.session.run(
            None,
            {
                "input_ids": input_ids,
                "attention_mask": attention_mask,
                "token_type_ids": token_type_ids,
            },
        )

        token_embeddings = outputs[0]
        pooled = self._mean_pooling(token_embeddings, attention_mask)
        return self._normalize(pooled)

    def embed_documents(self, texts: list[str], batch_size: int = 32) -> np.ndarray:
        """Embed a list of documents in batches."""
        all_embeddings = []
        for i in range(0, len(texts), batch_size):
            batch = texts[i:i + batch_size]
            all_embeddings.append(self.encode(batch))
        return np.vstack(all_embeddings)

    def embed_query(self, query: str) -> np.ndarray:
        """Embed a single query string."""
        return self.encode([query])[0]


class VectorStore:
    """Simple NumPy-based vector store with cosine similarity search."""

    def __init__(self, store_dir: Path = STORE_DIR):
        self.store_dir = store_dir
        self.embeddings: np.ndarray | None = None
        self.documents: list[dict] = []

    def save(self) -> None:
        """Persist store to disk."""
        self.store_dir.mkdir(parents=True, exist_ok=True)
        np.save(self.store_dir / "embeddings.npy", self.embeddings)
        with open(self.store_dir / "documents.json", "w", encoding="utf-8") as f:
            json.dump(self.documents, f, ensure_ascii=False, indent=2)

    def load(self) -> bool:
        """Load store from disk. Returns True if successful."""
        emb_path = self.store_dir / "embeddings.npy"
        doc_path = self.store_dir / "documents.json"
        if emb_path.exists() and doc_path.exists():
            self.embeddings = np.load(emb_path)
            with open(doc_path, "r", encoding="utf-8") as f:
                self.documents = json.load(f)
            return len(self.documents) > 0
        return False

    def add(self, texts: list[str], metadatas: list[dict], embeddings: np.ndarray) -> None:
        """Add documents with their embeddings."""
        self.documents = [
            {"text": t, "metadata": m} for t, m in zip(texts, metadatas)
        ]
        self.embeddings = embeddings

    def query(
        self,
        query_embedding: np.ndarray,
        top_k: int = 5,
        category: str | None = None,
    ) -> list[dict]:
        """Find top-k most similar documents using cosine similarity.

        Args:
            query_embedding: The query embedding vector.
            top_k: Maximum number of results to return.
            category: Optional routing category to restrict the search to
                (e.g. ``"policy"``). When ``None``, all documents are searched.

        Returns:
            A list of ``{"text", "source", "category", "score"}`` dicts ordered
            by similarity (highest first).
        """
        if self.embeddings is None or len(self.documents) == 0:
            return []

        # Already normalized, so dot product = cosine similarity.
        similarities = self.embeddings @ query_embedding

        # Candidate indices, optionally restricted to the requested category.
        if category is None:
            candidate_indices = np.argsort(similarities)[::-1]
        else:
            mask = [
                i for i in range(len(self.documents))
                if self.documents[i]["metadata"].get("category") == category
            ]
            if not mask:
                return []
            ordered = np.argsort(similarities[mask])[::-1]
            candidate_indices = [mask[i] for i in ordered]

        results = []
        for idx in candidate_indices[:top_k]:
            results.append({
                "text": self.documents[idx]["text"],
                "source": self.documents[idx]["metadata"]["source"],
                "category": self.documents[idx]["metadata"].get(
                    "category", "general"
                ),
                "score": float(similarities[idx]),
            })
        return results

    def count(self) -> int:
        """Return number of stored documents."""
        return len(self.documents)

    def clear(self) -> None:
        """Remove all documents and delete the persisted store files."""
        self.embeddings = None
        self.documents = []
        for name in ("embeddings.npy", "documents.json"):
            fpath = self.store_dir / name
            if fpath.exists():
                fpath.unlink()


class ChromaVectorStore:
    """Persistent vector store backed by ChromaDB.

    Provides the same interface as :class:`VectorStore` (``add`` / ``query`` /
    ``count`` / ``save`` / ``load``) so it is a drop-in replacement. Embeddings
    are supplied by the project's ONNX model, so ChromaDB is used purely as the
    vector index/persistence layer (no extra embedding dependency is pulled in).

    Vectors are persisted to disk under ``data/chroma_store/`` and reused across
    runs, avoiding a re-embed on every startup.
    """

    def __init__(self, persist_dir: Path = CHROMA_DIR) -> None:
        """Initialize the ChromaDB client and collection.

        Args:
            persist_dir: Directory where ChromaDB persists its data.

        Raises:
            ImportError: If the ``chromadb`` package is not installed.
        """
        try:
            import chromadb
            from chromadb.config import Settings
        except ImportError as exc:  # pragma: no cover - env-dependent
            raise ImportError(
                "The 'chromadb' package is required for VECTOR_BACKEND=chroma. "
                "Install it with 'pip install chromadb' or set "
                "VECTOR_BACKEND=numpy to use the legacy backend."
            ) from exc

        self.persist_dir = persist_dir
        self.persist_dir.mkdir(parents=True, exist_ok=True)
        self._client = chromadb.PersistentClient(
            path=str(self.persist_dir),
            settings=Settings(anonymized_telemetry=False, allow_reset=True),
        )
        # Cosine space matches the L2-normalized ONNX embeddings used elsewhere.
        self._collection = self._client.get_or_create_collection(
            name=CHROMA_COLLECTION,
            metadata={"hnsw:space": "cosine"},
        )

    def save(self) -> None:
        """No-op: ``PersistentClient`` writes to disk automatically."""
        return None

    def load(self) -> bool:
        """Report whether a persisted, non-empty collection already exists.

        Returns:
            True if the collection already contains vectors, else False.
        """
        return self._collection.count() > 0

    def add(
        self,
        texts: list[str],
        metadatas: list[dict],
        embeddings: np.ndarray,
    ) -> None:
        """Add documents and their precomputed embeddings to the collection.

        Args:
            texts: The chunk texts.
            metadatas: Per-chunk metadata dicts (must contain ``source``).
            embeddings: A ``(n, dim)`` array of L2-normalized embeddings.
        """
        ids = [f"chunk-{i}" for i in range(len(texts))]
        self._collection.add(
            ids=ids,
            documents=texts,
            metadatas=metadatas,
            embeddings=embeddings.tolist(),
        )

    def query(
        self,
        query_embedding: np.ndarray,
        top_k: int = 5,
        category: str | None = None,
    ) -> list[dict]:
        """Find the top-k most similar documents via ChromaDB.

        Args:
            query_embedding: The query embedding vector.
            top_k: Maximum number of results to return.
            category: Optional routing category to restrict the search to
                (e.g. ``"policy"``) using ChromaDB's native metadata ``where``
                filter. When ``None``, all documents are searched.

        Returns:
            A list of ``{"text", "source", "category", "score"}`` dicts ordered
            by similarity (highest first). ``score`` is cosine similarity in
            ``[-1, 1]`` to match the NumPy backend.
        """
        count = self._collection.count()
        if count == 0:
            return []

        where = {"category": category} if category is not None else None
        result = self._collection.query(
            query_embeddings=[query_embedding.tolist()],
            n_results=min(top_k, count),
            where=where,
            include=["documents", "metadatas", "distances"],
        )

        documents = result["documents"][0]
        metadatas = result["metadatas"][0]
        distances = result["distances"][0]

        results: list[dict] = []
        for text, metadata, distance in zip(documents, metadatas, distances):
            # Chroma returns cosine *distance* (1 - similarity); convert back
            # so downstream score thresholds behave like the NumPy backend.
            results.append({
                "text": text,
                "source": metadata.get("source", "unknown"),
                "category": metadata.get("category", "general"),
                "score": float(1.0 - distance),
            })
        return results

    def count(self) -> int:
        """Return the number of stored vectors.

        Returns:
            The number of documents in the collection.
        """
        return self._collection.count()

    def clear(self) -> None:
        """Delete and recreate the collection, removing all vectors."""
        self._client.delete_collection(CHROMA_COLLECTION)
        self._collection = self._client.get_or_create_collection(
            name=CHROMA_COLLECTION,
            metadata={"hnsw:space": "cosine"},
        )


def _make_store():
    """Create a vector store instance for the configured backend.

    Returns:
        A :class:`ChromaVectorStore` when ``VECTOR_BACKEND=chroma`` (default),
        otherwise a :class:`VectorStore` (NumPy backend).
    """
    if get_backend() == "chroma":
        return ChromaVectorStore()
    return VectorStore()


def load_documents() -> list[dict]:
    """Load documents recursively from the ``data/documents/`` folder.

    Walks the documents directory (including nested subfolders) and extracts
    text from every supported file. The relative path from the documents root
    is used as the ``source`` so files with the same name in different
    subfolders remain distinguishable.

    Supports: .txt, .md, .pdf, .docx, .pptx

    Returns:
        A list of ``{"content": str, "source": str}`` dicts, one per loaded
        document that yielded non-empty text.
    """
    if not DATA_DIR.exists():
        DATA_DIR.mkdir(parents=True, exist_ok=True)

    docs = []
    supported_extensions = {".txt", ".md", ".pdf", ".docx", ".pptx"}

    for fpath in sorted(DATA_DIR.rglob("*")):
        if not fpath.is_file():
            continue
        if fpath.suffix.lower() not in supported_extensions:
            continue

        source = fpath.relative_to(DATA_DIR).as_posix()
        try:
            text = _extract_text(fpath)
            if text.strip():
                docs.append({"content": text, "source": source})
                print(f"  📄 Loaded: {source} ({len(text)} chars)")
        except Exception as e:
            print(f"  ⚠️ Failed to load {source}: {e}")

    return docs


def _extract_text(fpath: Path) -> str:
    """Extract text content from a file based on its extension."""
    ext = fpath.suffix.lower()

    if ext in (".txt", ".md"):
        return fpath.read_text(encoding="utf-8")

    elif ext == ".pdf":
        try:
            import pymupdf  # PyMuPDF (modern import name)
        except ImportError:
            import fitz as pymupdf  # older PyMuPDF fallback
        text_parts = []
        with pymupdf.open(fpath) as doc:
            for page in doc:
                text_parts.append(page.get_text())
        return "\n\n".join(text_parts)

    elif ext == ".docx":
        from docx import Document
        doc = Document(fpath)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs)

    elif ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(fpath)
        text_parts = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        para_text = paragraph.text.strip()
                        if para_text:
                            slide_texts.append(para_text)
                # Extract text from tables in slides
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(
                            cell.text.strip() for cell in row.cells if cell.text.strip()
                        )
                        if row_text:
                            slide_texts.append(row_text)
            if slide_texts:
                text_parts.append(f"[Slide {slide_num}]\n" + "\n".join(slide_texts))
        return "\n\n".join(text_parts)

    else:
        return ""


def chunk_documents(docs: list[dict], chunk_size: int = 500, overlap: int = 100) -> list[dict]:
    """Split documents into chunks with metadata."""
    splitter = TextSplitter(chunk_size=chunk_size, chunk_overlap=overlap)
    chunks: list[dict] = []
    for doc in docs:
        splits = splitter.split_text(doc["content"])
        category = category_for_source(doc["source"])
        for i, text in enumerate(splits):
            chunks.append({
                "id": f"{doc['source']}_{i}",
                "text": text,
                "source": doc["source"],
                "category": category,
            })
    return chunks


def build_vector_store(chunks: list[dict], embeddings_model: EmbeddingModel):
    """Embed chunks and store them in the configured vector store.

    Args:
        chunks: The chunk dicts (``id``, ``text``, ``source``) to index.
        embeddings_model: The ONNX embedding model used to embed chunk text.

    Returns:
        The populated vector store (Chroma or NumPy, per ``VECTOR_BACKEND``).
    """
    texts = [c["text"] for c in chunks]
    metadatas = [
        {"source": c["source"], "category": c.get("category", "general")}
        for c in chunks
    ]
    embeddings = embeddings_model.embed_documents(texts)

    store = _make_store()
    store.add(texts, metadatas, embeddings)
    store.save()

    print(f"Indexed {store.count()} chunks into vector store "
          f"(backend: {get_backend()})")
    return store


def get_retriever() -> tuple[object, EmbeddingModel]:
    """Get or build the vector store for the configured backend.

    Loads a persisted store if one exists; otherwise ingests documents, chunks,
    embeds, and builds a fresh index.

    Returns:
        A ``(store, embedding_model)`` tuple. ``store`` is a
        :class:`ChromaVectorStore` or :class:`VectorStore` depending on
        ``VECTOR_BACKEND``.

    Raises:
        FileNotFoundError: If no supported documents exist when a build is
            required.
    """
    backend = get_backend()
    print(f"Vector backend: {backend}")

    store = _make_store()
    if store.load():
        return store, EmbeddingModel()

    print("📚 Building knowledge base from documents...")
    docs = load_documents()
    if not docs:
        raise FileNotFoundError(
            f"No supported documents found in {DATA_DIR}. "
            "Add .txt, .md, .pdf, .docx, or .pptx files (subfolders are "
            "searched recursively) to the data/documents/ folder."
        )
    chunks = chunk_documents(docs)
    embeddings_model = EmbeddingModel()
    store = build_vector_store(chunks, embeddings_model)
    return store, embeddings_model


def retrieve(
    query: str,
    store: object,
    embeddings_model: EmbeddingModel,
    top_k: int = 5,
    category: str | None = None,
) -> list[dict]:
    """Retrieve top-k relevant chunks for a query.

    Args:
        query: The user query text.
        store: The vector store to search (Chroma or NumPy backend).
        embeddings_model: The model used to embed the query.
        top_k: Maximum number of chunks to return.
        category: Optional routing category to restrict the search to
            (e.g. ``"policy"``). When ``None``, all documents are searched.

    Returns:
        A list of ``{"text", "source", "category", "score"}`` dicts ordered by
        similarity.
    """
    from src.observability import trace_span

    with trace_span("embedding") as attrs:
        query_embedding = embeddings_model.embed_query(query)
        attrs["query_len"] = len(query)
    return store.query(query_embedding, top_k=top_k, category=category)


# Admin uploads land here (a subfolder of the documents root, so they are
# picked up by the recursive loader). Files under "uploads/" map to the
# "general" category unless renamed into a known corpus folder.
UPLOADS_DIR = DATA_DIR / "uploads"

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".pptx"}


def save_uploaded_document(src_path: str, original_name: str | None = None) -> str:
    """Copy an uploaded file into the documents corpus.

    Args:
        src_path: Path to the uploaded temp file (as provided by the UI).
        original_name: Optional original filename to preserve. Falls back to
            the source file's name.

    Returns:
        The relative source label of the saved document (e.g.
        ``uploads/leave_policy.md``).

    Raises:
        ValueError: If the file extension is not supported.
    """
    src = Path(src_path)
    name = original_name or src.name
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(
            f"Unsupported file type '{ext}'. Supported types: {supported}."
        )

    UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    dest = UPLOADS_DIR / Path(name).name
    dest.write_bytes(src.read_bytes())
    return dest.relative_to(DATA_DIR).as_posix()


def rebuild_index() -> dict:
    """Clear and rebuild the vector index from all documents on disk.

    Re-ingests every supported document under ``data/documents/`` (including
    admin uploads), re-chunks, re-embeds, and repopulates the configured
    vector store. Used by the admin portal to apply new documents without a
    redeploy.

    Returns:
        A summary dict with ``backend``, ``documents``, and ``chunks`` counts.

    Raises:
        FileNotFoundError: If no supported documents are found.
    """
    from src.observability import log_event
    import logging

    backend = get_backend()
    store = _make_store()
    store.clear()

    docs = load_documents()
    if not docs:
        raise FileNotFoundError(
            f"No supported documents found in {DATA_DIR}. Add files before "
            "rebuilding the index."
        )
    chunks = chunk_documents(docs)
    embeddings_model = EmbeddingModel()
    store = build_vector_store(chunks, embeddings_model)

    summary = {
        "backend": backend,
        "documents": len(docs),
        "chunks": store.count(),
    }
    log_event(
        logging.INFO,
        f"index rebuilt: {summary['documents']} docs, "
        f"{summary['chunks']} chunks ({backend})",
        span="reindex",
        status="ok",
        **summary,
    )
    return summary
