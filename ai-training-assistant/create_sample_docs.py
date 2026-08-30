"""Create sample PDF, DOCX, and PPTX documents with dummy TechNova content."""

from pathlib import Path

DOCS_DIR = Path(__file__).parent / "data" / "documents"
DOCS_DIR.mkdir(parents=True, exist_ok=True)


def create_pdf():
    """Create a sample PDF: benefits.pdf"""
    import fitz  # pymupdf

    doc = fitz.open()

    # Page 1 - Benefits Overview
    page = doc.new_page()
    page.insert_text((72, 72), "TechNova Solutions — Employee Benefits Guide", fontsize=18)
    page.insert_text((72, 110), "=" * 60, fontsize=10)

    benefits_text = """
Health Insurance:
- Comprehensive medical coverage for employee + family (spouse & 2 children)
- Provider: Star Health Insurance
- Sum insured: ₹5,00,000 per year
- Cashless treatment at 5000+ network hospitals
- Pre-existing conditions covered after 2 years
- Dental and vision coverage included in premium plan

Life Insurance:
- Term life insurance: 3x annual CTC
- Accidental death benefit: 5x annual CTC
- Coverage starts from Day 1 of joining

Wellness Benefits:
- ₹5,000 annual wellness allowance (gym, yoga, sports)
- Free annual health checkup for employee and spouse
- Mental health support: 6 free counseling sessions per year
- Access to mindfulness app (Headspace subscription)

Retirement Benefits:
- Provident Fund (PF): 12% employer contribution
- Gratuity: As per Payment of Gratuity Act after 5 years
- National Pension Scheme (NPS): Optional, company matches up to 5%

Learning & Development:
- ₹50,000 annual learning budget per employee
- Access to Udemy Business, Coursera, and O'Reilly
- Conference attendance: 1 international or 2 domestic per year
- Internal certification programs with salary increments
"""
    y = 140
    for line in benefits_text.strip().split("\n"):
        page.insert_text((72, y), line, fontsize=10)
        y += 15
        if y > 750:
            page = doc.new_page()
            y = 72

    # Page 2 - Leave & Time Off
    page = doc.new_page()
    page.insert_text((72, 72), "Leave & Time Off Benefits", fontsize=16)
    page.insert_text((72, 100), "-" * 50, fontsize=10)

    leave_text = """
Paid Time Off:
- Annual Leave: 24 days per year (accrues at 2 days/month)
- Sick Leave: 12 days per year (medical certificate needed for >2 consecutive days)
- Casual Leave: 6 days per year (max 3 consecutive days)
- Maternity Leave: 26 weeks (as per Maternity Benefit Act)
- Paternity Leave: 10 working days
- Bereavement Leave: 5 days for immediate family

Special Leave:
- Marriage Leave: 5 days (once during employment)
- Exam Leave: 5 days per year (for approved certifications)
- Volunteer Leave: 2 days per year for community service
- Sabbatical: Up to 3 months unpaid after 3 years of service

Public Holidays:
- 12 public holidays per year (as per company holiday calendar)
- Floating holidays: 2 days (employee choice)

Work From Home:
- Hybrid model: 3 days office + 2 days WFH per week
- Full remote option available after 1 year of service (manager approval)
- WFH equipment allowance: ₹25,000 one-time setup
"""
    y = 120
    for line in leave_text.strip().split("\n"):
        page.insert_text((72, y), line, fontsize=10)
        y += 15

    output = DOCS_DIR / "benefits.pdf"
    doc.save(output)
    doc.close()
    print(f"✅ Created: {output.name}")


def create_docx():
    """Create a sample DOCX: it_setup_guide.docx"""
    from docx import Document

    doc = Document()
    doc.add_heading("TechNova Solutions — IT Setup Guide for New Employees", level=1)
    doc.add_paragraph("")

    doc.add_heading("Day 1: Account Setup", level=2)
    doc.add_paragraph(
        "On your first day, IT will provide you with a company laptop (MacBook Pro or Dell XPS based on role). "
        "Your accounts will be pre-configured with the following:"
    )
    items = [
        "Email: firstname.lastname@technova.com (Google Workspace)",
        "Slack: Auto-added to #general, #announcements, and your team channel",
        "GitHub: Enterprise account with access to team repositories",
        "Jira: Project management access based on your team assignment",
        "Confluence: Knowledge base and documentation access",
        "VPN: GlobalProtect VPN for secure remote access",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("VPN Setup Instructions", level=2)
    doc.add_paragraph(
        "To set up VPN access for remote work:"
    )
    steps = [
        "Download GlobalProtect from the IT portal (https://it.technova.internal/vpn)",
        "Install and open the application",
        "Enter portal address: vpn.technova.com",
        "Sign in with your company email and password",
        "Complete 2FA verification via Google Authenticator",
        "Once connected, you can access all internal systems remotely",
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f"{i}. {step}")

    doc.add_heading("Software & Tools", level=2)
    doc.add_paragraph("The following software is pre-installed on your company laptop:")

    # Add a table
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Tool"
    hdr_cells[1].text = "Purpose"
    hdr_cells[2].text = "Access"

    tools = [
        ("VS Code / IntelliJ", "Development IDE", "Pre-installed"),
        ("Docker Desktop", "Container development", "Pre-installed"),
        ("Postman", "API testing", "Pre-installed"),
        ("Figma", "Design collaboration", "Web-based, SSO login"),
        ("Notion", "Team documentation", "Web-based, SSO login"),
        ("Zoom", "Video conferencing", "Pre-installed"),
        ("1Password", "Password manager", "IT will send invite"),
    ]
    for tool, purpose, access in tools:
        row_cells = table.add_row().cells
        row_cells[0].text = tool
        row_cells[1].text = purpose
        row_cells[2].text = access

    doc.add_heading("IT Support & Troubleshooting", level=2)
    doc.add_paragraph(
        "For any IT issues, use the following channels:"
    )
    support = [
        "Slack: #it-helpdesk (response within 1 hour during business hours)",
        "Email: it.support@technova.com",
        "Phone: Extension 4500 (9 AM - 6 PM IST)",
        "Walk-in: IT desk on 3rd floor, Wing B",
        "Emergency (system down): Call +91-9876543210",
    ]
    for item in support:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Security Guidelines", level=2)
    doc.add_paragraph(
        "All employees must follow these security practices:"
    )
    security = [
        "Enable full-disk encryption (FileVault/BitLocker) — IT verifies on Day 1",
        "Use 1Password for all work-related passwords (never reuse personal passwords)",
        "Enable 2FA on all company accounts (Google, GitHub, AWS)",
        "Lock your screen when away from desk (Win+L or Cmd+Ctrl+Q)",
        "Never share credentials via Slack, email, or any messaging platform",
        "Report suspicious emails to security@technova.com immediately",
        "Connect only to TechNova-Secure WiFi in office (not TechNova-Guest)",
    ]
    for item in security:
        doc.add_paragraph(item, style="List Bullet")

    output = DOCS_DIR / "it_setup_guide.docx"
    doc.save(output)
    print(f"✅ Created: {output.name}")


def create_pptx():
    """Create a sample PPTX: company_culture.pptx"""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()

    # Slide 1 - Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "TechNova Solutions"
    slide.placeholders[1].text = "Company Culture & Values"

    # Slide 2 - Mission & Vision
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Our Mission & Vision"
    content = slide.placeholders[1]
    content.text = (
        "Mission: Empower businesses through intelligent technology solutions "
        "that drive growth, efficiency, and innovation.\n\n"
        "Vision: To be the leading AI-first technology company in Asia-Pacific "
        "by 2028, serving 10,000+ enterprise customers.\n\n"
        "Founded: 2015 in Bangalore, India\n"
        "Employees: 2,500+ across 8 offices\n"
        "Revenue: ₹800 Cr (FY 2025-26)"
    )

    # Slide 3 - Core Values
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Our 5 Core Values"
    content = slide.placeholders[1]
    content.text = (
        "1. Innovation First — We experiment boldly and learn from failures\n"
        "2. Customer Obsession — Every decision starts with the customer\n"
        "3. One Team — We collaborate across boundaries, no silos\n"
        "4. Ownership — We act like owners, not renters\n"
        "5. Continuous Learning — Growth mindset is non-negotiable"
    )

    # Slide 4 - Products
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Our Products"
    content = slide.placeholders[1]
    content.text = (
        "TechNova AI Platform — Enterprise AI/ML deployment platform\n"
        "DataFlow — Real-time data pipeline and analytics engine\n"
        "SecureGuard — AI-powered cybersecurity monitoring\n"
        "CloudBridge — Multi-cloud infrastructure management\n"
        "InsightBot — Conversational AI for customer support\n\n"
        "Key Clients: HDFC Bank, Infosys, Tata Motors, Flipkart, Apollo Hospitals"
    )

    # Slide 5 - Departments
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Departments & Teams"
    content = slide.placeholders[1]
    content.text = (
        "Engineering (60%) — Backend, Frontend, ML, DevOps, QA\n"
        "Product (10%) — Product Managers, Designers, Researchers\n"
        "Sales & Marketing (15%) — Enterprise Sales, Growth, Content\n"
        "People & Culture (5%) — HR, Talent Acquisition, L&D\n"
        "Finance & Operations (10%) — Finance, Legal, Admin, Facilities\n\n"
        "Reporting: Flat hierarchy, max 4 levels from CEO to IC"
    )

    # Slide 6 - Culture & Perks
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Culture & Perks"
    content = slide.placeholders[1]
    content.text = (
        "Work Model: Hybrid (3 office + 2 WFH days)\n"
        "Dress Code: Smart casual (no formals required)\n"
        "Team Events: Monthly team outings, quarterly offsites\n"
        "Hackathons: Quarterly 48-hour innovation sprints\n"
        "Friday Fun: Game rooms, foosball, board games after 4 PM\n"
        "Free Meals: Breakfast and lunch in office cafeteria\n"
        "Commute: Shuttle service from 15 pickup points\n"
        "Parking: Free for all employees (bike & car)"
    )

    # Slide 7 - Contact & Communication
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Communication Channels"
    content = slide.placeholders[1]
    content.text = (
        "Slack — Primary communication (async-first culture)\n"
        "Email — External communication and formal approvals\n"
        "Zoom — Video meetings (cameras on encouraged)\n"
        "Confluence — Documentation and knowledge sharing\n"
        "Town Halls — Monthly all-hands with CEO (every last Friday)\n"
        "Skip-levels — Quarterly 1:1 with skip-level manager\n\n"
        "Key Contacts:\n"
        "HR Helpdesk: hr.helpdesk@technova.com\n"
        "IT Support: it.support@technova.com / Slack #it-helpdesk\n"
        "Facilities: facilities@technova.com"
    )

    output = DOCS_DIR / "company_culture.pptx"
    prs.save(output)
    print(f"✅ Created: {output.name}")


if __name__ == "__main__":
    print("Creating sample documents in data/documents/...\n")
    create_pdf()
    create_docx()
    create_pptx()
    print("\n✅ All sample documents created!")
    print("   Delete data/vector_store/ and restart app.py to index them.")
